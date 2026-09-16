"""Tests for the xlsx writer, the Excel report, and the flat/HTML exporters."""

from __future__ import annotations

import csv
import datetime
import json
import re
import sys
import tempfile
import unittest
import zipfile
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree as ET

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from bomiq.config import Config  # noqa: E402
from bomiq.core.models import BomAnalysis  # noqa: E402
from bomiq.engine import Engine  # noqa: E402
from bomiq.export import flat, html as html_export, xlsx as xlsx_export  # noqa: E402
from bomiq.export.xlsx_writer import Workbook, cell_ref, col_letter  # noqa: E402
from bomiq.ingest.readers import read_xlsx_stdlib  # noqa: E402

SAMPLES = Path(__file__).resolve().parent.parent / "samples"


def temp_dir(prefix="bomiq-export-") -> Path:
    return Path(tempfile.mkdtemp(prefix=prefix))


def offline_analysis(sample="01-altium-export.csv", build=100) -> BomAnalysis:
    directory = temp_dir("bomiq-eng-")
    config = Config(data_dir=directory / "data",
                    config_dir=directory / "cfg", load_env=False,
                    use_keyring=False)
    config.settings.offline = True
    config.settings.build_quantity = build
    engine = Engine(config)
    try:
        return engine.analyse_file(SAMPLES / sample)
    finally:
        engine.close()


# --------------------------------------------------------------------------- #
# The xlsx writer
# --------------------------------------------------------------------------- #

class TestXlsxWriter(unittest.TestCase):
    def test_column_letters(self):
        self.assertEqual(col_letter(0), "A")
        self.assertEqual(col_letter(25), "Z")
        self.assertEqual(col_letter(26), "AA")
        self.assertEqual(col_letter(701), "ZZ")
        self.assertEqual(cell_ref(3, 2), "C3")

    def test_produces_a_valid_zip_with_the_expected_parts(self):
        wb = Workbook()
        sheet = wb.sheet("Data")
        sheet.row(["A", "B"])
        path = wb.save(temp_dir() / "x.xlsx")
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
        for required in ("[Content_Types].xml", "_rels/.rels",
                         "xl/workbook.xml", "xl/styles.xml",
                         "xl/sharedStrings.xml",
                         "xl/worksheets/sheet1.xml"):
            self.assertIn(required, names)

    def test_every_part_is_well_formed_xml(self):
        wb = Workbook()
        sheet = wb.sheet("Data", freeze="B2")
        sheet.row(["Header & <special>", "x"],
                  style=wb.style(bold=True, bg="FF0000"))
        sheet.row(["quote\"inside", 1.5])
        sheet.set_autofilter(1, 0, 2, 1)
        path = wb.save(temp_dir() / "y.xlsx")
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if name.endswith(".xml") or name.endswith(".rels"):
                    ET.fromstring(archive.read(name))

    def test_round_trip_values(self):
        wb = Workbook()
        sheet = wb.sheet("Types")
        sheet.row(["text", 42, 3.5, Decimal("0.0032"), True,
                   datetime.date(2026, 3, 4), ""])
        path = wb.save(temp_dir() / "z.xlsx")
        rows = read_xlsx_stdlib(path.read_bytes())[0].rows
        self.assertEqual(rows[0][0], "text")
        self.assertEqual(rows[0][1], "42")
        self.assertEqual(rows[0][2], "3.5")
        self.assertEqual(rows[0][3], "0.0032")

    def test_shared_strings_written_after_sheets(self):
        """Regression: strings registered while serialising must be included."""
        wb = Workbook()
        sheet = wb.sheet("S")
        sheet.row(["only-string-in-the-file"])
        path = wb.save(temp_dir() / "s.xlsx")
        with zipfile.ZipFile(path) as archive:
            shared = archive.read("xl/sharedStrings.xml").decode()
        self.assertIn("only-string-in-the-file", shared)

    def test_sheet_names_are_sanitised_and_unique(self):
        wb = Workbook()
        first = wb.sheet("A/B:C*D?E[F]G")
        second = wb.sheet("A-B-C-D-E-F-G")
        self.assertNotIn("/", first.name)
        self.assertNotEqual(first.name, second.name)

    def test_long_sheet_name_truncated(self):
        wb = Workbook()
        sheet = wb.sheet("x" * 60)
        self.assertLessEqual(len(sheet.name), 31)

    def test_illegal_control_characters_stripped(self):
        wb = Workbook()
        sheet = wb.sheet("S")
        sheet.row(["bad\x00char\x08here"])
        path = wb.save(temp_dir() / "c.xlsx")
        rows = read_xlsx_stdlib(path.read_bytes())[0].rows
        self.assertNotIn("\x00", rows[0][0])

    def test_hyperlinks_get_relationships(self):
        wb = Workbook()
        sheet = wb.sheet("L")
        sheet.row(["click"], hyperlinks=["https://example.com/a?b=1&c=2"])
        path = wb.save(temp_dir() / "l.xlsx")
        with zipfile.ZipFile(path) as archive:
            rels = archive.read(
                "xl/worksheets/_rels/sheet1.xml.rels").decode()
        self.assertIn("example.com", rels)
        ET.fromstring(rels)

    def test_gaps_preserve_column_alignment(self):
        wb = Workbook()
        sheet = wb.sheet("G")
        sheet.row(["a", "", "", "d"])
        path = wb.save(temp_dir() / "g.xlsx")
        rows = read_xlsx_stdlib(path.read_bytes())[0].rows
        self.assertEqual(rows[0][0], "a")
        self.assertEqual(rows[0][3], "d")


# --------------------------------------------------------------------------- #
# The Excel report
# --------------------------------------------------------------------------- #

class TestExcelReport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.analysis = offline_analysis()
        cls.path = xlsx_export.write_report(cls.analysis,
                                            temp_dir() / "report.xlsx")

    def test_all_sheets_present(self):
        sheets = [s.name for s in read_xlsx_stdlib(self.path.read_bytes())]
        for expected in ("Summary", "Enriched BOM", "Risk", "Sourcing",
                         "Alternates", "Compliance", "Issues", "Price curve",
                         "Providers"):
            self.assertIn(expected, sheets)

    def test_every_bom_line_appears(self):
        sheets = {s.name: s for s in read_xlsx_stdlib(self.path.read_bytes())}
        bom_sheet = sheets["Enriched BOM"]
        # header + one row per line
        self.assertEqual(bom_sheet.height, self.analysis.summary.total_lines + 1)

    def test_original_columns_are_echoed(self):
        sheets = {s.name: s for s in read_xlsx_stdlib(self.path.read_bytes())}
        header = sheets["Enriched BOM"].rows[0]
        self.assertTrue(any(cell.startswith("[src]") for cell in header))

    def test_offline_warning_is_in_the_summary(self):
        sheets = {s.name: s for s in read_xlsx_stdlib(self.path.read_bytes())}
        text = " ".join(cell for row in sheets["Summary"].rows
                        for cell in row)
        self.assertIn("OFFLINE MODE", text)

    def test_report_opens_in_openpyxl_when_available(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl is not installed")
        workbook = openpyxl.load_workbook(self.path)
        self.assertIn("Summary", workbook.sheetnames)
        grid = workbook["Enriched BOM"]
        self.assertEqual(grid.freeze_panes, "C2")
        self.assertTrue(grid.auto_filter.ref)


# --------------------------------------------------------------------------- #
# Flat exporters
# --------------------------------------------------------------------------- #

class TestFlatExports(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.analysis = offline_analysis()
        cls.dir = temp_dir()

    def test_csv_has_one_row_per_line(self):
        path = flat.write_csv(self.analysis, self.dir / "a.csv")
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        self.assertEqual(len(rows) - 1, self.analysis.summary.total_lines)
        self.assertIn("MPN", rows[0])

    def test_csv_is_excel_friendly(self):
        path = flat.write_csv(self.analysis, self.dir / "b.csv")
        self.assertTrue(path.read_bytes().startswith(b"\xef\xbb\xbf"))

    def test_issues_csv(self):
        path = flat.write_issues_csv(self.analysis, self.dir / "i.csv")
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        self.assertGreater(len(rows), 1)
        self.assertEqual(rows[0][0], "Severity")

    def test_quote_csv_excludes_dnp_and_scales_quantity(self):
        path = flat.write_quote_request_csv(self.analysis, self.dir / "q.csv")
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        self.assertTrue(rows)
        build = self.analysis.summary.build_quantity
        for row in rows:
            self.assertTrue(row["Manufacturer Part Number"])
            self.assertGreaterEqual(int(row["Quantity"]), build)

    def test_json_round_trips_through_the_model(self):
        path = flat.write_json(self.analysis, self.dir / "full.json",
                               full=True)
        reloaded = BomAnalysis.from_dict(json.loads(path.read_text()))
        self.assertEqual(len(reloaded.results), len(self.analysis.results))
        self.assertEqual(reloaded.health.score, self.analysis.health.score)
        self.assertEqual(reloaded.summary.total_cost,
                         self.analysis.summary.total_cost)

    def test_flat_json_is_a_list_of_records(self):
        path = flat.write_json(self.analysis, self.dir / "flat.json",
                               full=False)
        payload = json.loads(path.read_text())
        self.assertEqual(payload["format"], "bomiq.flat.v1")
        self.assertEqual(len(payload["lines"]),
                         self.analysis.summary.total_lines)
        self.assertIn("risk_score", payload["lines"][0])

    def test_summary_text_mentions_the_essentials(self):
        text = flat.summary_text(self.analysis)
        for token in ("Health", "Lines", "Cost", "Providers"):
            self.assertIn(token, text)

    def test_render_table_fits_the_width(self):
        table = flat.render_table(
            [["x" * 80, "y" * 80, "z" * 80]], ["a", "b", "c"], max_width=60)
        for line in table.splitlines():
            self.assertLessEqual(len(line), 62)


# --------------------------------------------------------------------------- #
# HTML report
# --------------------------------------------------------------------------- #

class TestHtmlReport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.analysis = offline_analysis()
        cls.html = html_export.render(cls.analysis)

    def test_is_self_contained(self):
        # No external resources at all.
        self.assertNotIn("<script src", self.html)
        self.assertNotIn("<link ", self.html)
        for pattern in ("src=\"http", "href=\"http://cdn",
                        "@import"):
            self.assertNotIn(pattern, self.html)

    def test_contains_every_line(self):
        for result in self.analysis.results:
            if result.line.mpn:
                self.assertIn(result.line.mpn, self.html)

    def test_has_the_health_score(self):
        self.assertIn(f"{self.analysis.health.score:.0f}", self.html)
        self.assertIn(self.analysis.health.grade, self.html)

    def test_offline_banner(self):
        self.assertIn("Offline mode", self.html)

    def test_html_is_escaped(self):
        analysis = self.analysis
        analysis.bom.name = '<script>alert("x")</script>'
        rendered = html_export.render(analysis)
        self.assertNotIn('<script>alert', rendered)
        self.assertIn("&lt;script&gt;", rendered)

    def test_writes_to_disk(self):
        path = html_export.write_html(self.analysis, temp_dir() / "r.html")
        self.assertGreater(path.stat().st_size, 5000)
        self.assertTrue(path.read_text(encoding="utf-8")
                        .startswith("<!DOCTYPE html>"))

    def test_no_unclosed_tag_obvious_breakage(self):
        opens = len(re.findall(r"<table", self.html))
        closes = len(re.findall(r"</table>", self.html))
        self.assertEqual(opens, closes)
        opens = len(re.findall(r"<tbody", self.html))
        closes = len(re.findall(r"</tbody>", self.html))
        self.assertEqual(opens, closes)


if __name__ == "__main__":
    unittest.main()
