"""BOM intelligence (SVG section 7): CSV/XLSX -> lifecycle distribution,
supply risk roll-up, cost roll-up, critical parts, alternate suggestions."""

from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass, field
from pathlib import Path

from app import schemas
from app.config import currency as default_currency
from app.normalizer import normalize_manufacturer_name
from app.service import get_part_report

_HEADERS = {
    "reference": (
        "reference",
        "ref",
        "refdes",
        "designator",
        "designators",
        "reference designator",
        "designation",
    ),
    "mpn": (
        "mpn",
        "partnumber",
        "part number",
        "part no",
        "part_no",
        "part",
        "manufacturer part number",
        "manufacturer p/n",
        "mfg part number",
        "manufacturer part",
        "p/n",
        "mfr part",
    ),
    "manufacturer": ("manufacturer", "mfg", "maker"),
    "qty": ("qty", "quantity", "count", "qty per board", "qty/b."),
    "target_price": ("target price", "price target", "max price", "target"),
    "value": ("value",),
}


def _header_match(header: str) -> str | None:
    key = header.strip().lower()
    for role, names in _HEADERS.items():
        if key in names:
            return role
    return None


def _read_csv(fh) -> list[dict]:
    return list(csv.DictReader(io.TextIOWrapper(fh, encoding="utf-8-sig", errors="replace")))


def _read_xlsx(fh) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(fh, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows)]
    out = []
    for r in rows:
        out.append({h: (v if v is not None else "") for h, v in zip(header, r)})
    return out


@dataclass
class BomLineResult:
    seq: int
    reference: str
    mpn: str
    manufacturer: str
    qty: int = 1
    target_price: float | None = None
    report: schemas.PartReportOut | None = None
    cost_estimate: float | None = None
    errors: list[str] = field(default_factory=list)

    @property
    def risk(self) -> schemas.RiskOut | None:
        return self.report.risk if self.report else None

    @property
    def critical(self) -> bool:
        return bool(self.risk and self.risk.level in ("HIGH", "CRITICAL"))

    @property
    def mpn_key(self) -> str:
        return self.mpn.upper()


@dataclass
class BomReport:
    filename: str
    lines: list[BomLineResult] = field(default_factory=list)
    currency: str = default_currency()

    @property
    def ok_lines(self) -> int:
        return sum(1 for ln in self.lines if ln.report is not None)

    @property
    def error_lines(self) -> int:
        return len(self.lines) - self.ok_lines

    def lifecycle_distribution(self) -> dict[str, int]:
        dist: dict[str, int] = {}
        for ln in self.lines:
            if ln.report is None:
                continue
            key = ln.report.component.lifecycle_status
            dist[key] = dist.get(key, 0) + 1
        return dist

    def risk_distribution(self) -> dict[str, int]:
        dist: dict[str, int] = {}
        for ln in self.lines:
            if ln.report is None:
                continue
            dist[ln.risk.level] = dist.get(ln.risk.level, 0) + 1
        return dist

    @property
    def critical_parts(self) -> list[BomLineResult]:
        return [ln for ln in self.lines if ln.critical]

    @property
    def total_cost_estimate(self) -> float | None:
        costs = [ln.cost_estimate for ln in self.lines if ln.cost_estimate is not None]
        return round(sum(costs), 2) if costs else None


def parse_bom_file(path: Path) -> list[dict]:
    if path.suffix.lower() in (".csv", ".txt"):
        with path.open("rb") as fh:
            return _read_csv(fh)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        with path.open("rb") as fh:
            return _read_xlsx(fh)
    raise ValueError(f"Unsupported BOM file type: {path.suffix}")


def _row_to_fields(row: dict) -> dict:
    role_map: dict[str, str] = {}
    for header, value in row.items():
        role = _header_match(header or "")
        if role and role not in role_map:
            role_map[role] = str(value or "").strip()
    return role_map


async def analyze_rows(rows: list[dict], progress=None) -> BomReport:
    report = BomReport(filename="upload")
    for i, row in enumerate(rows):
        fields = _row_to_fields(row)
        if not fields.get("mpn"):
            report.lines.append(
                BomLineResult(
                    seq=i + 1,
                    reference=fields.get("reference", ""),
                    mpn="",
                    manufacturer=fields.get("manufacturer", ""),
                    errors=["Missing MPN."],
                )
            )
            continue
        try:
            qty = int(float(fields["qty"])) if fields.get("qty") else 1
        except ValueError:
            qty = 1
        try:
            target_price = float(fields["target_price"]) if fields.get("target_price") else None
        except ValueError:
            target_price = None
        mfr = normalize_manufacturer_name(fields.get("manufacturer") or "")
        try:
            report_out = await get_part_report(fields["mpn"], mfr or None)
        except Exception as exc:  # noqa: BLE001 - one bad part must not kill the BOM
            report.lines.append(
                BomLineResult(
                    seq=i + 1,
                    reference=fields.get("reference", ""),
                    mpn=fields["mpn"],
                    manufacturer=fields.get("manufacturer", ""),
                    qty=qty,
                    target_price=target_price,
                    errors=[f"Lookup failed: {exc}"],
                )
            )
            continue
        resolved = bool(report_out.component.manufacturer) or bool(report_out.component.snapshots)
        if not resolved:
            report.lines.append(
                BomLineResult(
                    seq=i + 1,
                    reference=fields.get("reference", ""),
                    mpn=fields["mpn"],
                    manufacturer=fields.get("manufacturer", ""),
                    qty=qty,
                    target_price=target_price,
                    errors=["No source returned data for this MPN â€” verify the reference."],
                )
            )
            continue
        cost = None
        prices = []
        for snap in report_out.component.snapshots:
            for pb in snap.price_breaks:
                if pb.price:
                    prices.append(pb.price)
        if prices:
            cost = round(min(prices) * qty, 2)
        line = BomLineResult(
            seq=i + 1,
            reference=fields.get("reference", ""),
            mpn=fields["mpn"],
            manufacturer=fields.get("manufacturer", ""),
            qty=qty,
            target_price=target_price,
            report=report_out,
            cost_estimate=cost,
        )
        if cost and target_price and cost > target_price:
            line.errors.append(f"Estimated cost {cost} exceeds target {target_price}.")
        report.lines.append(line)
        if progress:
            progress(i + 1, len(rows))
    report.filename = f"bom_{len(report.lines)}items"
    return report


def to_bom_report_out(report: BomReport, filename: str) -> schemas.BomReportOut:
    def line_out(ln: BomLineResult) -> schemas.BomLine:
        return schemas.BomLine(
            seq=ln.seq,
            reference=ln.reference,
            mpn=ln.mpn,
            manufacturer=ln.manufacturer,
            qty=ln.qty,
            target_price=ln.target_price,
            component=ln.report.component if ln.report else None,
            risk=ln.report.risk if ln.report else None,
            alternatives=ln.report.alternatives if ln.report else [],
            cost_estimate=ln.cost_estimate,
            errors=ln.errors,
        )

    return schemas.BomReportOut(
        filename=filename,
        total_lines=len(report.lines),
        ok_lines=report.ok_lines,
        error_lines=report.error_lines,
        lifecycle_distribution=report.lifecycle_distribution(),
        risk_distribution=report.risk_distribution(),
        critical_parts=[line_out(ln) for ln in report.critical_parts],
        total_cost_estimate=report.total_cost_estimate,
        currency=report.currency,
        lines=[line_out(ln) for ln in report.lines],
    )


def export_bom_json(report: schemas.BomReportOut) -> str:
    return json.dumps(report.model_dump(), indent=2, default=str)


def export_bom_csv(report: schemas.BomReportOut) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(
        [
            "seq",
            "reference",
            "mpn",
            "manufacturer",
            "qty",
            "lifecycle",
            "risk",
            "risk_score",
            "n_vendors",
            "min_price",
            "cost_estimate",
            "alternatives",
            "errors",
        ]
    )
    for line in report.lines:
        n_vendors = len({s.distributor for s in line.component.snapshots}) if line.component else 0
        prices = []
        if line.component:
            for snap in line.component.snapshots:
                for pb in snap.price_breaks:
                    if pb.price:
                        prices.append(pb.price)
        min_price = min(prices) if prices else ""
        alts = "; ".join(f"{a.mpn}({a.kind})" for a in line.alternatives[:3]) if line.alternatives else ""
        writer.writerow(
            [
                line.seq,
                line.reference,
                line.mpn,
                line.manufacturer,
                line.qty,
                line.component.lifecycle_status if line.component else "",
                line.risk.level if line.risk else "",
                line.risk.score if line.risk else "",
                n_vendors,
                min_price,
                line.cost_estimate or "",
                alts,
                "; ".join(line.errors),
            ]
        )
    return buf.getvalue()
